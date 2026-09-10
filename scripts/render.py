#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""render.py — IBL 交付物渲染器（样式固定来自 references/design-spec.md、samples 样例与模板版式）

用法:
  python render.py docx       <blocks|lesson_spec|course_spec>.json <output.docx>
  python render.py rubric     <lesson_spec.json> <output.docx>
  python render.py pptx       <ppt_brief.json> <template.pptx> <output.pptx>
  python render.py mktemplate <template.pptx> <output.pptx>

pptx 依赖模板中的 IBL_* 版式（封面/论点/大图/讨论/编号行/竖卡/三栏）。
模板为原始设计页版本时，先运行一次 mktemplate 生成版式；模板更换后重跑。

输入 JSON 支持 "-" 从 stdin 读入。解析内置容错：未转义的 ASCII 引号、
字符串内裸换行、缺失逗号、末尾逗号会被自动修复。
"""
import argparse
import json
import re
import sys

# ---- 设计令牌（references/design-spec.md）----
PRIMARY = "0083FE"      # 品牌主色蓝
PRIMARY_200 = "99CFFF"  # 浅蓝描边
PRIMARY_50 = "E6F3FF"   # 最浅蓝底（表头/占位框）
CYAN_600 = "00CCC0"     # 荧光青悬浮态（渐变终点，保证浅字可读）
DARK = "1A1A1A"         # 标题重色
TEXT = "333333"         # 正文
MUTED = "808080"        # 弱化信息
BORDER = "D9D9D9"       # 边框/分隔线
FONT = "Noto Sans SC"   # 与 samples 样例一致（思源黑体）

# ---- docx 字号（pt，对齐 samples/前测-AI科学研究入门.docx）----
H1, H2, BODY, NOTE, TBL = 18, 10, 9.5, 7.5, 9


# ============================================================
# JSON 读取与容错修复
# ============================================================
def repair_json(text):
    """修复 LLM 常见 JSON 笔误：串内裸 ASCII 引号、裸换行/Tab、缺失/末尾逗号。

    仅在严格解析失败后调用，不影响合法 JSON。
    """
    out = []
    i, n = 0, len(text)
    in_str = False
    sig = ""  # 最近一个非空白输出字符
    while i < n:
        c = text[i]
        if not in_str:
            if c in '"{[' and sig and sig in '"}]':
                out.append(",")  # 值之间缺失逗号
                sig = ","
            if c == '"':
                in_str = True
            out.append(c)
            if not c.isspace():
                sig = c
            i += 1
            continue
        if c == "\\" and i + 1 < n:
            out.append(text[i:i + 2])
            sig = ""
            i += 2
            continue
        if c == '"':
            j = i + 1
            while j < n and text[j].isspace():
                j += 1
            nxt = text[j] if j < n else ""
            if nxt in ",:}]{[":  # 结构性后随字符 → 视为字符串结束
                in_str = False
                out.append(c)
                sig = c
            else:  # 否则为内容引号 → 转义
                out.append('\\"')
            i += 1
            continue
        if c in "\n\r\t":
            out.append({"\n": "\\n", "\r": "\\r", "\t": "\\t"}[c])
            i += 1
            continue
        out.append(c)
        if not c.isspace():
            sig = c
        i += 1
    return re.sub(r",\s*([}\]])", r"\1", "".join(out))


def load_json(path):
    if path == "-":
        raw = sys.stdin.read()
    else:
        with open(path, encoding="utf-8-sig") as f:
            raw = f.read()
    try:
        return json.loads(raw)
    except json.JSONDecodeError as first_err:
        try:
            return json.loads(repair_json(raw))
        except json.JSONDecodeError:
            raise ValueError("JSON 解析失败（自动修复后仍无效）：%s" % first_err)


# ============================================================
# DOCX：通用块渲染
# ============================================================
def new_docx():
    from docx import Document
    from docx.shared import Pt, Mm, Cm, RGBColor
    from docx.oxml.ns import qn

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(4)
    return doc


def write_blocks(doc, blocks):
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def fmt(run, size, bold, color):
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)

    def para(text, size=BODY, bold=False, color=TEXT,
             space_before=0, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        for i, seg in enumerate(str(text).split("\n")):
            if i:
                p.add_run().add_break()
            r = p.add_run(seg)
            fmt(r, size, bold, color)
        return p

    def borders(table):
        tblPr = table._tbl.tblPr
        el = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement("w:" + edge)
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), BORDER)
            el.append(e)
        look = tblPr.find(qn("w:tblLook"))
        if look is not None:
            tblPr.insert(list(tblPr).index(look), el)
        else:
            tblPr.append(el)

    def shade(cell, fill):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shd)

    def cell_text(cell, text, bold, color):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(str(text))
        fmt(r, TBL, bold, color)

    def table_block(spec):
        header = spec.get("header")
        rows = spec.get("rows", [])
        ncol = len(header) if header else (len(rows[0]) if rows else 0)
        if not ncol:
            return
        table = doc.add_table(rows=(1 if header else 0) + len(rows), cols=ncol)
        borders(table)
        ri = 0
        if header:
            for ci, val in enumerate(header):
                c = table.rows[0].cells[ci]
                cell_text(c, val, True, PRIMARY)
                shade(c, PRIMARY_50)
            ri = 1
        for row in rows:
            for ci, val in enumerate(row[:ncol]):
                cell_text(table.rows[ri].cells[ci], val, False, TEXT)
            ri += 1
        widths = spec.get("widths")
        if widths:
            table.autofit = False
            for ci, w in enumerate(widths[:ncol]):
                for row in table.rows:
                    row.cells[ci].width = Cm(float(w))
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    for b in blocks:
        if "h1" in b:
            para(b["h1"], H1, True, PRIMARY, space_before=6, space_after=10)
        elif "h2" in b:
            para(b["h2"], H2, True, DARK, space_before=10)
        elif "p" in b:
            para(b["p"])
        elif "note" in b:
            para(b["note"], NOTE, False, MUTED)
        elif "answer_lines" in b:
            for _ in range(int(b["answer_lines"])):
                para("＿" * 46, BODY, False, TEXT, space_after=2)
        elif "check" in b:
            para("☐ " + b["check"])
        elif "table" in b:
            table_block(b["table"])


# ============================================================
# DOCX：spec 直读构建器（SKILL.md 7.2 映射规则的代码化）
# ============================================================
def anchor_blocks(anchor):
    """research_anchor → note 块序列：关键词一行，来源逐条列出"""
    if not anchor:
        return []
    blocks = []
    keywords = anchor.get("keywords") or []
    if keywords:
        blocks.append({"note": "研究锚点 · 检索关键词：" + " / ".join(keywords)})
    for src in anchor.get("sources") or []:
        title = str(src.get("title") or "").strip()
        if not title:
            continue
        text = "· " + title
        if src.get("year"):
            text += f" ({src['year']})"
        if src.get("url"):
            text += f" {src['url']}"
        blocks.append({"note": text})
    return blocks


def outline_blocks(spec):
    """course_spec → 课程大纲块序列"""
    b = []
    theme = spec.get("theme", "课程")
    b.append({"h1": f"无界实验室 · {theme} · 课程大纲"})

    positioning = spec.get("positioning")
    if not positioning:
        ft = spec.get("final_task") or {}
        positioning = (f"本课程以“{theme}”为主题，围绕驱动问题“{spec.get('driving_question', '')}”"
                       f"展开系列真实探究。学生将经历完整的“提问—设计—实证—解释—表达”链路，"
                       f"最终完成：{ft.get('deliverable', '综合探究成果')}。")
    b.append({"h2": "一 课程定位"})
    b.append({"p": positioning})

    aud = spec.get("audience") or {}
    parts = [f"年级/年龄：{aud.get('grade', '待定')}"]
    if aud.get("class_size"):
        parts.append(f"班级人数：{aud['class_size']}")
    if aud.get("grouping"):
        parts.append(f"分组方式：{aud['grouping']}")
    b.append({"h2": "二 学生画像"})
    b.append({"p": "\n".join(parts)})

    b.append({"h2": "三 学习目标"})
    objs = spec.get("objectives") or []
    if objs:
        b.append({"table": {"header": ["维度", "目标"],
                            "rows": [[o.get("dimension", ""), o.get("objective", "")] for o in objs],
                            "widths": [3, 13.5]}})

    b.append({"h2": "四 课程驱动问题"})
    b.append({"p": spec.get("driving_question", "")})
    b.extend(anchor_blocks(spec.get("research_anchor")))

    b.append({"h2": "五 课程模块"})
    for m in spec.get("modules", []):
        b.append({"p": f"{m.get('id', '')} {m.get('title', '')} —— 子问题：{m.get('sub_question', '')}"})
        rows = [[l.get("lesson", ""), l.get("title", ""),
                 l.get("research_question", ""), l.get("student_task", "")]
                for l in m.get("lessons", [])]
        if rows:
            b.append({"table": {"header": ["课次", "课题", "研究问题", "学生任务"],
                                "rows": rows, "widths": [1.5, 3.5, 6, 5.5]}})

    b.append({"h2": "六 进阶路径"})
    dur = spec.get("duration") or {}
    ft = spec.get("final_task") or {}
    total = dur.get("total_lessons", 0)
    if total:
        b.append({"p": f"全课程共 {total} 课，每课 {dur.get('minutes_per_lesson', '—')} 分钟。"
                       f"前约 75% 课时按模块递进探究，从现象观察与提问走向方案设计、实验取证与解释论证。"})
    else:
        b.append({"p": "前约 75% 课时按模块递进探究，从现象观察与提问走向方案设计、实验取证与解释论证。"})
    if ft:
        b.append({"p": f"最后约 25% 课时（第 {ft.get('lesson_range', '—')} 课）：综合应用小闭环。"
                       f"挑战：{ft.get('challenge', '')}；最终成果：{ft.get('deliverable', '')}。"})
        applied = ft.get("applied_knowledge") or []
        if applied:
            b.append({"p": "综合运用：" + "；".join(applied) + "。"})
        loop = ft.get("closed_loop") or []
        if loop:
            b.append({"p": "闭环环节：" + " → ".join(loop)})
    return b


def plan_blocks(spec):
    """lesson_spec → 教案块序列"""
    b = []
    les = spec.get("lesson") or {}
    b.append({"h1": f"无界实验室 · 教案 · {les.get('title', '单节课')}"})
    no = f"第 {les['lesson_no']} 课" if les.get("lesson_no") else "—"
    b.append({"table": {"header": ["模块", "课次", "时长", "已有基础"],
                        "rows": [[les.get("module", "—"), no,
                                  f"{les.get('duration_min', '—')} 分钟",
                                  les.get("prior_learning", "—")]],
                        "widths": [4, 2.5, 2.5, 7.5]}})

    b.append({"h2": "一 研究问题"})
    b.append({"p": spec.get("research_question", "")})
    b.extend(anchor_blocks(spec.get("research_anchor")))

    b.append({"h2": "二 核心学生任务"})
    ct = spec.get("core_task") or {}
    b.append({"p": ct.get("task", "")})
    choices = ct.get("student_choices") or []
    if choices:
        b.append({"p": "学生自主选择空间：" + "；".join(choices) + "。"})
    if ct.get("expected_output"):
        b.append({"p": "预期产出：" + ct["expected_output"]})

    b.append({"h2": "三 节点安排"})
    rows = []
    total = 0
    for i, f in enumerate(spec.get("flow", []), 1):
        mins = f.get("minutes", 0)
        total += mins
        rows.append([f"节点{i}", f"{mins} 分钟",
                     f.get("student_activity", ""), f.get("teacher_support", "")])
    rows.append(["合计", f"{total} 分钟", "", ""])
    b.append({"table": {"header": ["节点", "时长", "学生活动", "教师支持"],
                        "rows": rows, "widths": [1.8, 2, 7, 5.7]}})

    cards = spec.get("evidence_cards") or []
    if cards:
        b.append({"h2": "四 证据卡"})
        b.append({"p": "、".join(cards)})

    res = spec.get("resources") or {}
    diff = spec.get("differentiation") or {}
    if any(res.get(k) for k in ("per_group", "shared", "teacher_prep", "media")) or \
            diff.get("support") or diff.get("challenge"):
        b.append({"h2": "五 资源与差异化"})
        if res.get("per_group"):
            b.append({"p": "每组材料：" + "；".join(res["per_group"])})
        if res.get("shared"):
            b.append({"p": "共用设备：" + "；".join(res["shared"])})
        if res.get("teacher_prep"):
            b.append({"p": "教师课前准备：" + "；".join(res["teacher_prep"])})
        if res.get("media"):
            b.append({"p": "媒体资源：" + "；".join(res["media"])})
        if diff.get("support"):
            b.append({"p": "基础支持：" + diff["support"]})
        if diff.get("challenge"):
            b.append({"p": "拓展挑战：" + diff["challenge"]})
    return b


def rubric_blocks(spec):
    """lesson_spec → 行为观察评价块序列"""
    les = spec.get("lesson") or {}
    b = [{"h1": f"无界实验室 · 行为观察评价 · {les.get('title', '单节课')}"}]
    info = [f"课程/模块：{les.get('module', '—')}",
            f"课次：{les.get('lesson_no', '—')}",
            "班级：____________    日期：____________"]
    b.append({"p": "\n".join(info)})
    b.append({"h2": "一 观察记录（看到对应行为即勾选）"})
    for item in spec.get("assessment", []):
        b.append({"check": item})
    b.append({"h2": "二 学生勾选区"})
    b.append({"table": {"header": ["学生姓名", "勾选记录", "备注（选填）"],
                        "rows": [["", "", ""]] * 8, "widths": [3.5, 6, 7]}})
    b.append({"h2": "三 备注"})
    b.append({"answer_lines": 2})
    return b


def to_blocks(data):
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        if "blocks" in data:
            return data["blocks"]
        if "modules" in data:
            return outline_blocks(data)
        if "flow" in data:
            return plan_blocks(data)
    raise ValueError("无法识别的输入格式：应为 blocks / course_spec（含 modules）/ lesson_spec（含 flow）")


def cmd_docx(input_path, out_path):
    blocks = to_blocks(load_json(input_path))
    doc = new_docx()
    write_blocks(doc, blocks)
    doc.save(out_path)


def cmd_rubric(input_path, out_path):
    doc = new_docx()
    write_blocks(doc, rubric_blocks(load_json(input_path)))
    doc.save(out_path)


# ============================================================
# PPTX：模板设计页 → IBL_* 版式（mktemplate），再按版式渲染（pptx）
# ============================================================
IBL_LAYOUTS = {
    "cover": "IBL_COVER",
    "points": "IBL_POINTS",
    "image": "IBL_IMAGE",
    "quote": "IBL_QUOTE",
    "rows": "IBL_ROWS",
    "cards": "IBL_CARDS",
    "columns": "IBL_COLUMNS",
}

# 模板设计页槽位表：(规范化文本, x, y, 占位符类型)，坐标单位英寸、允许 ±0.15 误差
IBL_DESIGNS = [
    ("IBL_COVER", {
        0: ("标题", 2.70, 2.76, "ctrTitle"),
        1: ("副标题", 2.73, 4.02, "subTitle"),
    }),
    ("IBL_POINTS", {
        0: ("章节标题", 0.71, 0.28, "title"),
        1: ("论点标题", 0.98, 1.61, "body"),
        2: ("论点详情", 0.98, 2.43, "body"),
        3: ("论点标题", 5.35, 1.24, "body"),
        4: ("论点1", 5.58, 2.06, "body"),
        5: ("论点详情", 5.58, 2.43, "body"),
        6: ("论点2", 5.58, 3.54, "body"),
        7: ("论点详情", 5.58, 3.91, "body"),
        8: ("论点3", 5.58, 5.02, "body"),
        9: ("论点详情", 5.58, 5.39, "body"),
    }),
    ("IBL_IMAGE", {
        0: ("章节标题", 0.71, 0.28, "title"),
        1: ("图片", 1.26, 1.43, "pic"),
    }),
    ("IBL_QUOTE", {
        0: ("章节标题", 0.71, 0.28, "title"),
        1: ("图片", 3.66, 1.25, "pic"),
        2: ("", 0.70, 1.25, "body"),
        3: ("", 3.66, 4.05, "body"),
        4: ("", 7.20, 1.25, "body"),
    }),
    ("IBL_ROWS", {
        0: ("章节标题", 0.71, 0.28, "title"),
        1: ("论点", 1.05, 1.62, "body"),
        2: ("论点详情", 1.05, 2.42, "body"),
        3: ("论点", 6.65, 1.22, "body"),
        4: ("分论点1", 7.85, 1.93, "body"),
        5: ("分论点1详情", 7.85, 2.30, "body"),
        6: ("分论点2", 7.85, 3.48, "body"),
        7: ("分论点2详情", 7.85, 3.85, "body"),
        8: ("分论点3", 7.85, 5.03, "body"),
        9: ("分论点3详情", 7.85, 5.40, "body"),
    }),
    ("IBL_CARDS", {
        0: ("章节标题", 0.71, 0.28, "title"),
        1: ("一", 2.19, 2.23, "body"),
        2: ("二", 6.34, 2.23, "body"),
        3: ("三", 10.49, 2.23, "body"),
        4: ("论点一", 1.10, 3.00, "body"),
        5: ("论点二", 5.25, 3.00, "body"),
        6: ("论点三", 9.40, 3.00, "body"),
    }),
    ("IBL_COLUMNS", {
        0: ("章节标题", 0.71, 0.28, "title"),
        1: ("论点", 1.05, 1.86, "body"),
        2: ("一", 1.70, 3.28, "body"),
        3: ("副标题", 1.70, 3.66, "body"),
        4: ("详情", 1.18, 4.14, "body"),
        5: ("二", 5.75, 3.28, "body"),
        6: ("副标题", 5.75, 3.66, "body"),
        7: ("详情", 5.23, 4.14, "body"),
        8: ("三", 9.80, 3.28, "body"),
        9: ("副标题", 9.80, 3.66, "body"),
        10: ("详情", 9.28, 4.14, "body"),
    }),
]

# 原设计缺失、需新建的占位符：(idx, x, y, w, h, 对齐, 字号, 颜色)
IBL_NEW_SLOTS = {
    "IBL_IMAGE": [(2, 1.26, 6.32, 10.56, 0.40, "ctr", 1200, "595959")],
    "IBL_CARDS": [(7, 1.00, 3.74, 3.15, 2.36, "l", 1100, "404040"),
                  (8, 5.15, 3.74, 3.15, 2.36, "l", 1100, "404040"),
                  (9, 9.30, 3.74, 3.15, 2.36, "l", 1100, "404040")],
}


def cmd_mktemplate(src_path, out_path):
    import copy
    import zipfile
    from lxml import etree

    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    CT = "http://schemas.openxmlformats.org/package/2006/content-types"
    REL = "http://schemas.openxmlformats.org/package/2006/relationships"
    LAYOUT_CT = "application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"
    EMU = 914400
    TOL = int(0.15 * EMU)

    def E(tag, ns):
        return "{%s}%s" % (ns, tag)

    zin = zipfile.ZipFile(src_path)
    files = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    names = set(files)

    def root(name):
        return etree.fromstring(files[name])

    def dump(tree):
        return etree.tostring(tree, xml_declaration=True,
                              encoding="UTF-8", standalone=True)

    # ---- 定位幻灯片与母版 ----
    pres = root("ppt/presentation.xml")
    prels = root("ppt/_rels/presentation.xml.rels")
    rel_tgt = {rel.get("Id"): rel.get("Target") for rel in prels}
    slide_parts = []
    for sld in pres.find(E("sldIdLst", P)):
        t = rel_tgt[sld.get(E("id", R))].lstrip("/")
        slide_parts.append(t if t.startswith("ppt/") else "ppt/" + t)
    if len(slide_parts) < len(IBL_DESIGNS):
        raise ValueError("模板仅有 %d 页，不足 %d 张设计页" %
                         (len(slide_parts), len(IBL_DESIGNS)))
    slide_parts = slide_parts[:len(IBL_DESIGNS)]

    master_part = "ppt/slideMasters/slideMaster1.xml"
    master_rels_part = "ppt/slideMasters/_rels/slideMaster1.xml.rels"
    master = root(master_part)
    master_rels = root(master_rels_part)
    next_id = 2147483647
    for n in names:
        if re.match(r"ppt/slideMasters/slideMaster\d+\.xml$", n):
            for lid in root(n).findall(".//" + E("sldLayoutId", P)):
                next_id = max(next_id, int(lid.get("id")))
    next_id += 1
    next_rid = 1 + max([int(r.get("Id")[3:]) for r in master_rels
                        if (r.get("Id") or "").startswith("rId")] + [0])
    next_lay = 1
    while "ppt/slideLayouts/slideLayout%d.xml" % next_lay in names:
        next_lay += 1
    ct = root("[Content_Types].xml")
    slide_w = int(pres.find(".//" + E("sldSz", P)).get("cx"))

    def shape_geo(sp):
        off = sp.find(".//" + E("off", A))
        ext = sp.find(".//" + E("ext", A))
        return (int(off.get("x")), int(off.get("y")),
                int(ext.get("cx")), int(ext.get("cy")))

    def norm_text(sp):
        if sp.tag != E("sp", P):
            return ""
        txt = "".join(t.text or "" for t in sp.findall(".//" + E("t", A)))
        txt = re.sub(r"\s+", "", txt)
        return re.sub(r"^\d*章节标题$", "章节标题", txt)

    def inject_style(sp):
        """把槽位形状首段的段落/字符格式写进版式占位符 lstStyle，供幻灯片继承"""
        tx = sp.find(E("txBody", P))
        if tx is None:
            return
        p0 = tx.find(E("p", A))
        if p0 is None:
            return
        lst = tx.find(E("lstStyle", A))
        if lst is None:
            lst = etree.Element(E("lstStyle", A))
            tx.insert(list(tx).index(p0), lst)
        for ch in list(lst):
            lst.remove(ch)
        lvl1 = etree.SubElement(lst, E("lvl1pPr", A))
        pPr = p0.find(E("pPr", A))
        if pPr is not None:
            for k, v in pPr.attrib.items():
                lvl1.set(k, v)
            for ch in pPr:
                if ch.tag != E("defRPr", A):
                    lvl1.append(copy.deepcopy(ch))
        bullet_tags = {E("buNone", A), E("buChar", A), E("buAutoNum", A)}
        if not any(ch.tag in bullet_tags for ch in lvl1):
            pos = next((i for i, ch in enumerate(lvl1)
                        if ch.tag in (E("tabLst", A), E("extLst", A))), len(lvl1))
            lvl1.insert(pos, etree.Element(E("buNone", A)))
        r0 = p0.find(E("r", A))
        rPr = r0.find(E("rPr", A)) if r0 is not None else None
        if rPr is None:
            rPr = p0.find(E("endParaRPr", A))
        if rPr is not None:
            d = copy.deepcopy(rPr)
            d.tag = E("defRPr", A)
            lvl1.append(d)

    def promote(sp, idx, phtype):
        if phtype == "pic" and sp.tag == E("sp", P):
            sp.tag = E("pic", P)
            nv = sp.find(E("nvSpPr", P))
            nv.tag = E("nvPicPr", P)
            cnv = nv.find(E("cNvSpPr", P))
            if cnv is not None:
                i = list(nv).index(cnv)
                nv.remove(cnv)
                nv.insert(i, etree.Element(E("cNvPicPr", P)))
            tx = sp.find(E("txBody", P))
            if tx is not None:
                sp.remove(tx)
            sp.insert(1, etree.Element(E("blipFill", P)))
        else:
            inject_style(sp)
        nvPr = sp.find(".//" + E("nvPr", P))
        ph = etree.Element(E("ph", P))
        ph.set("type", phtype)
        ph.set("idx", str(idx))
        nvPr.insert(0, ph)

    def new_text_ph(idx, x, y, w, h, align, size_pt, color):
        anchor = "ctr" if align == "ctr" else "t"
        algn = ' algn="%s"' % align if align != "l" else ""
        xml = (
            '<p:sp xmlns:p="%s" xmlns:a="%s">'
            '<p:nvSpPr><p:cNvPr id="%d" name="IBL Slot %d"/>'
            '<p:cNvSpPr txBox="1"/><p:nvPr><p:ph type="body" idx="%d"/></p:nvPr></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            '<p:txBody><a:bodyPr wrap="square" anchor="%s"><a:normAutofit/></a:bodyPr>'
            "<a:lstStyle><a:lvl1pPr%s><a:buNone/><a:defRPr sz=\"%d\"><a:solidFill>"
            '<a:srgbClr val="%s"/></a:solidFill></a:defRPr></a:lvl1pPr></a:lstStyle>'
            '<a:p><a:pPr%s/><a:endParaRPr lang="zh-CN"/></a:p></p:txBody></p:sp>'
        ) % (P, A, 900 + idx, idx, idx,
             int(x * EMU), int(y * EMU), int(w * EMU), int(h * EMU),
             anchor, algn, size_pt, color, algn)
        return etree.fromstring(xml.encode("utf-8"))

    removed_parts = set(slide_parts)

    for di, (layout_name, slots) in enumerate(IBL_DESIGNS):
        part = slide_parts[di]
        sld = root(part)
        spTree = sld.find(E("cSld", P) + "/" + E("spTree", P))

        for sp in list(spTree):
            if sp.tag in (E("sp", P), E("pic", P)):
                x, y, w, h = shape_geo(sp)
                if x >= slide_w:  # 画面外残留装饰
                    spTree.remove(sp)

        cands = []
        for sp in spTree:
            if sp.tag not in (E("sp", P), E("pic", P)):
                continue
            x, y, w, h = shape_geo(sp)
            cands.append((norm_text(sp), x, y, sp))

        used = set()
        for idx in sorted(slots):
            want, wx, wy, phtype = slots[idx]
            hit = None
            for c in cands:
                if id(c[3]) in used:
                    continue
                if c[0] == want and abs(c[1] - wx * EMU) <= TOL \
                        and abs(c[2] - wy * EMU) <= TOL:
                    hit = c
                    break
            if hit is None:
                raise ValueError("%s：槽位 %d（%s @ %.2f,%.2f）未在模板第 %d 页找到，"
                                 "模板与预期设计页不符" %
                                 (layout_name, idx, want or "空白面板", wx, wy, di + 1))
            used.add(id(hit[3]))
            promote(hit[3], idx, phtype)

        for spec in IBL_NEW_SLOTS.get(layout_name, []):
            spTree.append(new_text_ph(*spec))

        sld.tag = E("sldLayout", P)
        for tag in ("transition", "timing"):
            el = sld.find(E(tag, P))
            if el is not None:
                sld.remove(el)
        sld.find(E("cSld", P)).set("name", layout_name)

        rels_part = part.replace("ppt/slides/", "ppt/slides/_rels/") + ".rels"
        new_rels = etree.Element(E("Relationships", REL), nsmap={None: REL})
        if rels_part in names:
            for rel in root(rels_part):
                ty = rel.get("Type", "")
                if ty.endswith("/slideLayout") or "notesSlide" in ty:
                    continue
                new_rels.append(copy.deepcopy(rel))
        m_rel = etree.SubElement(new_rels, E("Relationship", REL))
        m_rel.set("Id", "rIdM%d" % di)
        m_rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster")
        m_rel.set("Target", "../slideMasters/slideMaster1.xml")

        layout_part = "ppt/slideLayouts/slideLayout%d.xml" % next_lay
        files[layout_part] = dump(sld)
        files["ppt/slideLayouts/_rels/slideLayout%d.xml.rels" % next_lay] = dump(new_rels)
        ov = etree.SubElement(ct, E("Override", CT))
        ov.set("PartName", "/" + layout_part)
        ov.set("ContentType", LAYOUT_CT)

        rel_node = etree.SubElement(master_rels, E("Relationship", REL))
        rel_node.set("Id", "rId%d" % next_rid)
        rel_node.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout")
        rel_node.set("Target", "../slideLayouts/slideLayout%d.xml" % next_lay)
        lid = etree.SubElement(master.find(".//" + E("sldLayoutIdLst", P)),
                               E("sldLayoutId", P))
        lid.set("id", str(next_id))
        lid.set(E("id", R), "rId%d" % next_rid)
        next_id += 1
        next_rid += 1
        next_lay += 1

        rels_part_path = rels_part
        if rels_part_path in names:
            removed_parts.add(rels_part_path)
            for rel in root(rels_part_path):
                if "notesSlide" in (rel.get("Type") or ""):
                    t = (rel.get("Target") or "").replace("../", "")
                    removed_parts.add("ppt/" + t)

    # ---- 删除原幻灯片 ----
    sldlst = pres.find(E("sldIdLst", P))
    for sld in list(sldlst):
        sldlst.remove(sld)
    for rel in list(prels):
        if (rel.get("Type") or "").endswith("/slide"):
            prels.remove(rel)
    root_ext = pres.find(E("extLst", P))
    if root_ext is not None and any(
            etree.QName(ch).localname == "sectionLst" for ch in root_ext.iter()):
        pres.remove(root_ext)

    files["ppt/presentation.xml"] = dump(pres)
    files["ppt/_rels/presentation.xml.rels"] = dump(prels)
    files[master_part] = dump(master)
    files[master_rels_part] = dump(master_rels)
    files["[Content_Types].xml"] = dump(ct)

    for ovn in list(ct):
        if ovn.tag == E("Override", CT):
            pn = (ovn.get("PartName") or "").lstrip("/")
            if pn in removed_parts:
                ct.remove(ovn)
    files["[Content_Types].xml"] = dump(ct)

    zout = zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED)
    for n, data in files.items():
        if n in removed_parts:
            continue
        zout.writestr(n, data)
    zout.close()

    from pptx import Presentation
    prs = Presentation(out_path)
    cover = next(ly for m in prs.slide_masters for ly in m.slide_layouts
                 if ly.name == "IBL_COVER")
    prs.slides.add_slide(cover)
    prs.save(out_path)
    print("OK %s (IBL layouts: %d)" % (out_path, len(IBL_DESIGNS)))


# ---- ppt_brief 语义 → 内容池 ----
def _pt(p, key):
    return str((p or {}).get(key) or "").strip()


def _norm_points(s):
    ps = s.get("points")
    if ps:
        return [p if isinstance(p, dict) else {"detail": str(p)} for p in ps]
    c = str(s.get("content") or "").strip()
    if not c:
        return []
    segs = [x.strip() for x in re.split(r"[\n。！？；;]+", c) if x.strip()]
    return [{"detail": x} for x in segs]


def _jt(p):
    return "\n".join(x for x in (_pt(p, "title"), _pt(p, "detail")) if x)


def _choose_layout(s, i, alt):
    lay = str(s.get("layout") or "").strip().lower()
    if lay in IBL_LAYOUTS:
        return lay
    if (s.get("media") or {}).get("type") in ("image", "video"):
        return "image"
    pts = _norm_points(s)
    if i == 0:
        return "cover"
    if len(pts) == 2:
        return "quote"
    if len(pts) == 3:
        return "columns" if any(_pt(p, "detail") for p in pts) else "cards"
    if len(pts) >= 4:
        alt[0] += 1
        return "points" if alt[0] % 2 else "rows"
    return "points"


def _fill_map(layout_name, s, pts, brief):
    topic = str(s.get("topic") or "").strip()
    sub = str(s.get("subtitle") or "").strip()
    media = s.get("media") or {}

    def pt(i, key):
        return _pt(pts[i] if i < len(pts) else None, key)

    def jt(i):
        return _jt(pts[i] if i < len(pts) else None)

    if layout_name == "IBL_COVER":
        return {0: topic or str(brief.get("lesson_spec_ref") or "课程课件"),
                1: sub or str(brief.get("subtitle") or "")}
    if layout_name == "IBL_POINTS" or layout_name == "IBL_ROWS":
        m = {0: topic, 3: sub}
        for k, base in enumerate((1, 4, 6, 8)):
            m[base] = pt(k, "title")
            m[base + 1] = pt(k, "detail")
        return m
    if layout_name == "IBL_IMAGE":
        return {0: topic,
                2: str(media.get("description") or "") or sub}
    if layout_name == "IBL_QUOTE":
        return {0: topic, 2: jt(0), 3: sub, 4: jt(1)}
    if layout_name == "IBL_CARDS":
        m = {0: topic}
        for i in range(3):
            m[1 + i] = pt(i, "badge") or "%02d" % (i + 1)
            m[4 + i] = pt(i, "title")
            m[7 + i] = pt(i, "detail")
        return m
    m = {0: topic, 1: sub}
    for i in range(3):
        m[2 + 3 * i] = pt(i, "title")
        m[4 + 3 * i] = pt(i, "detail")
    return m


def render_pptx(brief_path, template_path, out_path):
    from pptx import Presentation
    from pptx.oxml.ns import qn

    brief = load_json(brief_path)
    prs = Presentation(template_path)
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.get(qn("r:id")))
        sldIdLst.remove(sldId)
    layouts = {ly.name: ly for m in prs.slide_masters for ly in m.slide_layouts}
    missing = [n for n in IBL_LAYOUTS.values() if n not in layouts]
    if missing:
        raise ValueError("模板缺少版式 %s。请先运行："
                         "python scripts/render.py mktemplate <模板> <输出>" %
                         "、".join(missing))

    slides = brief.get("slides") or []
    if not slides:
        raise ValueError("ppt_brief.slides 为空")

    alt = [0]
    for i, s in enumerate(slides):
        key = _choose_layout(s, i, alt)
        slide = prs.slides.add_slide(layouts[IBL_LAYOUTS[key]])
        texts = _fill_map(IBL_LAYOUTS[key], s, _norm_points(s), brief)
        idxs = {}
        for ph in slide.placeholders:
            idxs[ph.placeholder_format.idx] = ph
            if ph.has_text_frame:
                ph.text_frame.clear()
        for idx, text in texts.items():
            ph = idxs.get(idx)
            if ph is None or not ph.has_text_frame or not text:
                continue
            tf = ph.text_frame
            for j, seg in enumerate(str(text).split("\n")):
                para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                run = para.add_run()
                run.text = seg
        note = str(s.get("teacher_action") or "").strip()
        if note:
            slide.notes_slide.notes_text_frame.text = note

    prs.save(out_path)


def main():
    ap = argparse.ArgumentParser(description="IBL 交付物渲染器")
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("docx", help="blocks / course_spec / lesson_spec → DOCX（自动路由）")
    d.add_argument("input")
    d.add_argument("output")
    r = sub.add_parser("rubric", help="lesson_spec → 行为观察评价 DOCX")
    r.add_argument("input")
    r.add_argument("output")
    p = sub.add_parser("pptx", help="ppt_brief JSON + 模板 → PPTX")
    p.add_argument("brief")
    p.add_argument("template")
    p.add_argument("output")
    k = sub.add_parser("mktemplate", help="模板设计页 → IBL_* 版式（一次性）")
    k.add_argument("src")
    k.add_argument("output")
    a = ap.parse_args()
    if a.cmd == "docx":
        cmd_docx(a.input, a.output)
    elif a.cmd == "rubric":
        cmd_rubric(a.input, a.output)
    elif a.cmd == "mktemplate":
        cmd_mktemplate(a.src, a.output)
    else:
        render_pptx(a.brief, a.template, a.output)
    print("OK %s" % a.output)


if __name__ == "__main__":
    main()
